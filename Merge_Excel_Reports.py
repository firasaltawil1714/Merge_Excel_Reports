from pathlib import Path
from openpyxl import Workbook, load_workbook


class ExcelMerger:
    def __init__(self, reports_folder, output_path):
        self.reports_folder = reports_folder
        self.output_path = output_path
        self.all_rows = []

    def get_excel_files(self):
        return [f for f in self.reports_folder.iterdir() if f.suffix == ".xlsx"]

    def read_rows(self, file_path):
        workbook = load_workbook(file_path)
        sheet = workbook.active
        return [tuple(cell.value for cell in row) for row in sheet.iter_rows()]

    def merge(self):
        files = self.get_excel_files()
        all_rows = []
        for i, file in enumerate(files):
            rows = self.read_rows(file)
            if i == 0:
                all_rows.extend(rows)
            else:
                all_rows.extend(rows[1:])
        self.all_rows = all_rows

    def save(self):
        workbook = Workbook()
        sheet = workbook.active
        for row in self.all_rows:
            sheet.append(row)
        workbook.save(self.output_path)

    def run(self):
        files = self.get_excel_files()
        if not files:
            print("No Excel files found.")
            return
        self.merge()
        self.save()
        print(f"Merged {len(files)} files into {self.output_path.name}")


if __name__ == "__main__":
    script_dir = Path(__file__).parent
    merger = ExcelMerger(reports_folder=script_dir / "reports", output_path=script_dir / "merged_report.xlsx")
    merger.run()
